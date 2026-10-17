"""Glue between the Streamlit grids and the inventory model. No Streamlit imports, so it is testable."""

import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ghg_inventory.excel_io import (
    ACTIVITY_COLUMNS, METRICS, REFRIGERANT_COLUMNS, make_template, parse_tables,
)
from ghg_inventory.models import ActivityData, Config, Factor

FRAME_NAMES = ("activity", "refrigerants", "metrics")
_SHEETS = {"activity": "Activity", "refrigerants": "Refrigerants", "metrics": "Metrics"}

_ACTIVITY_DTYPES = {"source_id": "object", "month": "object", "quantity": "float64", "unit": "object",
                    "data_quality": "Int64", "evidence_ref": "object", "notes": "object"}
_REFRIGERANT_DTYPES = {"equipment": "object", "refrigerant": "object", "charge_kg": "float64", "method": "object",
                       "annual_leak_rate": "float64", "recharged_kg": "float64", "data_quality": "Int64",
                       "evidence_ref": "object"}


def _typed(df: pd.DataFrame, dtypes: dict[str, str]) -> pd.DataFrame:
    df = df.reindex(columns=list(dtypes))
    for column, dtype in dtypes.items():
        if dtype == "object":
            df[column] = df[column].astype("object").where(df[column].notna(), None)
        else:
            df[column] = pd.to_numeric(df[column], errors="coerce").astype(dtype)
    return df.reset_index(drop=True)


def metrics_frame(values: dict[str, float]) -> pd.DataFrame:
    return pd.DataFrame(
        [{"metric": key, "value": values.get(key), "unit": unit, "notes": note} for key, unit, note in METRICS]
    ).astype({"value": "float64"})


def frames_from_data(data: ActivityData) -> dict[str, pd.DataFrame]:
    activity = pd.DataFrame([[a.source_id, a.month, a.quantity, a.unit, a.data_quality, a.evidence_ref, a.notes]
                             for a in data.activities], columns=list(ACTIVITY_COLUMNS))
    refrigerants = pd.DataFrame([[r.equipment, r.refrigerant, r.charge_kg, r.method, r.annual_leak_rate,
                                  r.recharged_kg, r.data_quality, r.evidence_ref] for r in data.refrigerants],
                                columns=list(REFRIGERANT_COLUMNS))
    return {
        "activity": _typed(activity, _ACTIVITY_DTYPES),
        "refrigerants": _typed(refrigerants, _REFRIGERANT_DTYPES),
        "metrics": metrics_frame(data.metrics),
    }


def empty_frames() -> dict[str, pd.DataFrame]:
    return frames_from_data(ActivityData(activities=[]))


def frames_from_workbook(file) -> dict[str, pd.DataFrame]:
    """Read a workbook's sheets into grids without validating, so rows with errors can be fixed in the UI."""
    wb = load_workbook(file, data_only=True, read_only=True)
    missing = [name for name in _SHEETS.values() if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"The workbook has no {', '.join(missing)} sheet. Start from the input template.")

    def rows(sheet, width):
        out = []
        for row in wb[sheet].iter_rows(min_row=2, max_col=width, values_only=True):
            values = list(row) + [None] * (width - len(row))
            if any(v not in (None, "") for v in values):
                out.append(values)
        return out

    activity = pd.DataFrame(rows("Activity", len(ACTIVITY_COLUMNS)), columns=list(ACTIVITY_COLUMNS))
    activity["month"] = activity["month"].map(
        lambda m: m.strftime("%Y-%m") if hasattr(m, "strftime") else m)
    refrigerants = pd.DataFrame(rows("Refrigerants", len(REFRIGERANT_COLUMNS)), columns=list(REFRIGERANT_COLUMNS))
    metric_values = {}
    for key, value, *_ in rows("Metrics", 2):
        try:
            metric_values[key] = float(value)
        except (TypeError, ValueError):
            pass
    return {
        "activity": _typed(activity, _ACTIVITY_DTYPES),
        "refrigerants": _typed(refrigerants, _REFRIGERANT_DTYPES),
        "metrics": metrics_frame(metric_values),
    }


def _rows(df: pd.DataFrame, columns: tuple[str, ...]) -> list[tuple[int, list]]:
    out = []
    for i, record in enumerate(df.reindex(columns=list(columns)).itertuples(index=False), start=1):
        values = []
        for v in record:
            if v is pd.NA or (isinstance(v, float) and pd.isna(v)):
                v = None
            elif hasattr(v, "item"):          # numpy scalar -> python
                v = v.item()
            values.append(v)
        out.append((i, values))
    return out


def data_from_frames(frames: dict[str, pd.DataFrame], config: Config) -> tuple[ActivityData | None, list[str]]:
    """Validate the grids. Problems name the grid row, e.g. 'Activity row 3, quantity: ...'."""
    metric_rows = [(i, [m, v]) for i, (m, v) in enumerate(
        zip(frames["metrics"]["metric"], frames["metrics"]["value"]), start=1)]
    metric_rows = [(i, [m, None if pd.isna(v) else float(v)]) for i, (m, v) in metric_rows]
    return parse_tables(
        config,
        _rows(frames["activity"], ACTIVITY_COLUMNS),
        _rows(frames["refrigerants"], REFRIGERANT_COLUMNS),
        metric_rows,
        sheets=("Activity", "Refrigerants", "Metrics"),
        cell_refs=False,
    )


def backup(path: Path) -> Path | None:
    if not path.exists():
        return None
    target = path.with_name(f"{path.stem}.{datetime.now():%Y%m%d-%H%M%S}.bak{path.suffix}")
    shutil.copy2(path, target)
    return target


def save_workbook(config: Config, factors: dict[str, Factor], data: ActivityData, path: Path) -> Path | None:
    """Write validated data to the input workbook (template layout). Returns the backup of the old file."""
    previous = backup(path)
    make_template(config, factors, path, data=data)
    return previous


def frames_equal(a: dict[str, pd.DataFrame], b: dict[str, pd.DataFrame]) -> bool:
    return all(a[name].reset_index(drop=True).equals(b[name].reset_index(drop=True)) for name in FRAME_NAMES)


def plausibility_warnings(activity: pd.DataFrame) -> list[str]:
    """Things that pass validation but look wrong: a month far above every other month of the same
    source (a typo or a missing decimal point), and rows that repeat the same source, month and amount."""
    warnings = []
    df = activity.dropna(subset=["source_id", "month", "quantity"]).copy()
    df["row"] = df.index + 1
    for (source, unit), group in df.groupby(["source_id", "unit"], dropna=True):
        if len(group) < 4:
            continue
        ordered = group.sort_values("quantity", ascending=False)
        top, runner_up = ordered.iloc[0], ordered.iloc[1]
        if runner_up["quantity"] > 0 and top["quantity"] > 2.5 * runner_up["quantity"]:
            warnings.append(
                f"Activity row {int(top['row'])}: {top['quantity']:,.0f} {unit} of {source} in {top['month']} is "
                f"{top['quantity'] / runner_up['quantity']:,.1f} times the next highest month. Check for a typo.")
    duplicated = df[df.duplicated(subset=["source_id", "month", "quantity"], keep=False)]
    for _, group in duplicated.groupby(["source_id", "month", "quantity"]):
        rows = ", ".join(str(int(r)) for r in group["row"])
        first = group.iloc[0]
        warnings.append(f"Activity rows {rows} repeat {first['source_id']} for {first['month']} with the same "
                        "amount. Remove one if it is the same bill entered twice.")
    return warnings

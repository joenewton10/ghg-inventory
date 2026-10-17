"""Excel input template: generate it, and read a filled one back with cell-level error messages.

A folder holding activity.csv, refrigerants.csv and metrics.csv (same columns) is accepted too.
"""

import csv
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError

from ghg_inventory.factors import UNIT_CHOICES
from ghg_inventory.gwp import REFRIGERANT_CHOICES
from ghg_inventory.models import DATA_QUALITY, ActivityData, ActivityRecord, Config, Factor, RefrigerantRecord

ACTIVITY_COLUMNS = ("source_id", "month", "quantity", "unit", "data_quality", "evidence_ref", "notes")
REFRIGERANT_COLUMNS = ("equipment", "refrigerant", "charge_kg", "method", "annual_leak_rate",
                       "recharged_kg", "data_quality", "evidence_ref")
METRIC_COLUMNS = ("metric", "value", "unit", "notes")
METRICS = (
    ("material_processed_t", "tonne", "Material cut and converted in the year (intensity denominator)"),
    ("material_purchased_t", "tonne", "Incoming filter media bought in the year (Scope 3 screening)"),
    ("fte", "FTE", "Average full-time equivalent employees"),
    ("floor_area_m2", "m2", "Heated floor area of the site"),
    ("revenue_eur", "EUR", "Net revenue (VSME / ESRS intensity denominator)"),
)

_EDITABLE_ROWS = 600
_HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


class InputError(ValueError):
    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("input data has problems:\n" + "\n".join(f"  {p}" for p in problems))


# ---------------------------------------------------------------- template


def _sheet(wb: Workbook, title: str, columns: tuple[str, ...], widths: dict[str, int]):
    ws = wb.create_sheet(title)
    ws.append(list(columns))
    for idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 16)
    ws.freeze_panes = "A2"
    return ws


def _protect_header(ws, n_columns: int):
    """Lock the header row only; data cells stay editable."""
    for row in ws.iter_rows(min_row=2, max_row=_EDITABLE_ROWS, max_col=n_columns):
        for cell in row:
            cell.protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.formatRows = False
    ws.protection.insertRows = False
    ws.protection.deleteRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


def _list_validation(ws, column: str, formula: str, prompt: str):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True)
    dv.error, dv.errorTitle, dv.prompt = f"Pick a value from the list. {prompt}", "Not in list", prompt
    ws.add_data_validation(dv)
    dv.add(f"{column}2:{column}{_EDITABLE_ROWS}")


def make_template(config: Config, factors: dict[str, Factor], path: str | Path,
                  data: ActivityData | None = None) -> Path:
    """Write the input workbook. With data, the rows are filled in (used for the representative file)."""
    path = Path(path)
    wb = Workbook()
    intro = wb.active
    intro.title = "Instructions"
    year = config.reporting.year
    lines = [
        f"GHG inventory input, {config.organisation.name}, reporting year {year}",
        "",
        "Fill one row per invoice, meter reading or fuel-card statement in the Activity sheet.",
        "source_id: pick from the list. The Sources sheet says what each one covers and which unit it expects.",
        "month: the month the energy was used (for a bill covering several months, split it or use the end month).",
        "unit: kWh, MWh, litre or m3. Gas in m3 is converted with Zustandszahl x Brennwert from config.yaml.",
        "data_quality: 1 to 5, see the scale below. Be honest; it drives the uncertainty range.",
        "evidence_ref: invoice number or file name, so an auditor can trace the row.",
        "",
        "Refrigerants: one row per piece of equipment that holds refrigerant (AC, chillers, refrigerant air dryers).",
        "  screening method: give charge_kg (nameplate) and annual_leak_rate (e.g. 0.05 for 5%).",
        "  mass_balance method: give recharged_kg, the refrigerant topped up during the year (from service records).",
        "",
        "Metrics: denominators for intensity figures and the Scope 3 screen. Leave blank if unknown.",
        "",
        "Data quality scale:",
        *[f"  {k}: {label} (treated as ±{u:g}% in the uncertainty analysis)" for k, (label, u) in DATA_QUALITY.items()],
    ]
    for line in lines:
        intro.append([line])
    intro["A1"].font = Font(bold=True, size=13)
    intro.column_dimensions["A"].width = 120

    lists = wb.create_sheet("Lists")
    months = [f"{year}-{m:02d}" for m in range(1, 13)]
    columns = {
        "A": [s.id for s in config.sources],
        "B": months,
        "C": list(UNIT_CHOICES),
        "D": list(REFRIGERANT_CHOICES),
        "E": ["screening", "mass_balance"],
    }
    for col, values in columns.items():
        for i, value in enumerate(values, start=1):
            lists[f"{col}{i}"] = value
    lists.sheet_state = "hidden"

    activity = _sheet(wb, "Activity", ACTIVITY_COLUMNS, {"source_id": 20, "evidence_ref": 28, "notes": 40})
    _list_validation(activity, "A", f"Lists!$A$1:$A${len(columns['A'])}", "Source from the Sources sheet")
    _list_validation(activity, "B", "Lists!$B$1:$B$12", f"Month in {year}")
    _list_validation(activity, "D", f"Lists!$C$1:$C${len(UNIT_CHOICES)}", "Unit")
    dq = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True)
    dq.error = "Data quality is a whole number from 1 (metered) to 5 (rough estimate)."
    activity.add_data_validation(dq)
    dq.add(f"E2:E{_EDITABLE_ROWS}")

    refrigerants = _sheet(wb, "Refrigerants", REFRIGERANT_COLUMNS, {"equipment": 32, "evidence_ref": 28})
    _list_validation(refrigerants, "B", f"Lists!$D$1:$D${len(REFRIGERANT_CHOICES)}", "Refrigerant type (nameplate)")
    _list_validation(refrigerants, "D", "Lists!$E$1:$E$2", "screening or mass_balance")
    dq_r = DataValidation(type="whole", operator="between", formula1="1", formula2="5", showErrorMessage=True)
    refrigerants.add_data_validation(dq_r)
    dq_r.add(f"G2:G{_EDITABLE_ROWS}")

    metrics = _sheet(wb, "Metrics", METRIC_COLUMNS, {"metric": 24, "notes": 70})
    values = (data.metrics if data else {})
    for key, unit, note in METRICS:
        metrics.append([key, values.get(key), unit, note])

    sources = _sheet(wb, "Sources", ("source_id", "name", "scope", "category", "factor", "expected_unit"),
                     {"name": 36, "category": 24, "factor": 36})
    for s in config.sources:
        factor_id = s.factor or s.location_factor
        sources.append([s.id, s.name, s.scope, s.category, factor_id, factors[factor_id].unit])
    sources.protection.sheet = True

    if data:
        for r in data.activities:
            activity.append([r.source_id, r.month, r.quantity, r.unit, r.data_quality, r.evidence_ref, r.notes])
        for r in data.refrigerants:
            refrigerants.append([r.equipment, r.refrigerant, r.charge_kg, r.method, r.annual_leak_rate,
                                 r.recharged_kg, r.data_quality, r.evidence_ref])

    for ws in (activity, refrigerants):
        _protect_header(ws, ws.max_column)
    _protect_header(metrics, 2)
    for ws in (activity, refrigerants, metrics, sources):
        for cell in ws[1]:
            cell.alignment = Alignment(vertical="center")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- reading


def _blank(value) -> bool:
    return _clean(value) is None


def _clean(value):
    """Strip strings; treat empty strings and NaN (from pandas) as missing."""
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _where(sheet: str, row_number: int, columns: tuple[str, ...], field: str | None, cell_refs: bool) -> str:
    """'Activity!C7 (quantity)' for workbooks, 'Activity row 7, quantity' for the UI grids."""
    if cell_refs:
        if field in columns:
            return f"{sheet}!{get_column_letter(columns.index(field) + 1)}{row_number} ({field})"
        return f"{sheet}!{row_number}"
    return f"{sheet} row {row_number}" + (f", {field}" if field else "")


def _parse(sheet: str, columns: tuple[str, ...], rows, model, problems: list[str], cell_refs: bool = True) -> list:
    records = []
    for row_number, values in rows:
        values = list(values)[: len(columns)]
        if all(_blank(v) for v in values):
            continue
        raw = {name: _clean(v) for name, v in zip(columns, values)}
        if "month" in raw and raw["month"] is not None and not isinstance(raw["month"], str):
            raw["month"] = raw["month"].strftime("%Y-%m") if hasattr(raw["month"], "strftime") else str(raw["month"])
        raw = {k: v for k, v in raw.items() if v is not None}
        for text_field in ("evidence_ref", "notes"):
            if text_field in raw:
                raw[text_field] = str(raw[text_field])
        try:
            row_ref = f"{sheet}!{row_number}" if cell_refs else f"{sheet} row {row_number}"
            records.append(model(row_ref=row_ref, **raw))
        except ValidationError as exc:
            for err in exc.errors():
                field = err["loc"][0] if err["loc"] else None
                where = _where(sheet, row_number, columns, field if field in columns else None, cell_refs)
                problems.append(f"{where}: {err['msg']}")
    return records


def _parse_metrics(sheet: str, rows, problems: list[str], cell_refs: bool = True) -> dict[str, float]:
    known = {key for key, _, _ in METRICS}
    metrics = {}
    for row_number, values in rows:
        values = list(values) + [None, None]
        key, value = _clean(values[0]), _clean(values[1])
        if key is None or value is None:
            continue
        if key not in known:
            problems.append(f"{_where(sheet, row_number, METRIC_COLUMNS, 'metric', cell_refs)}: unknown metric {key!r}")
            continue
        try:
            metrics[key] = float(value)
        except (TypeError, ValueError):
            problems.append(f"{_where(sheet, row_number, METRIC_COLUMNS, 'value', cell_refs)}: "
                            f"{key} must be a number, got {value!r}")
    return metrics


def _check_sources(config: Config, records: list[ActivityRecord], problems: list[str], cell_refs: bool = True):
    ids = {s.id for s in config.sources}
    year = str(config.reporting.year)
    for r in records:
        def where(field):
            return f"{r.row_ref} ({field})" if cell_refs else f"{r.row_ref}, {field}"
        if r.source_id not in ids:
            problems.append(f"{where('source_id')}: {r.source_id!r} is not in the source register in config.yaml")
        if not r.month.startswith(year):
            problems.append(f"{where('month')}: {r.month} is outside the reporting year {year}")


def parse_tables(config: Config, activity_rows, refrigerant_rows, metric_rows,
                 sheets: tuple[str, str, str] = ("Activity", "Refrigerants", "Metrics"),
                 cell_refs: bool = True) -> tuple[ActivityData | None, list[str]]:
    """Validate rows of (row_number, values). Returns (data, problems); data is None if there are problems."""
    problems: list[str] = []
    activities = _parse(sheets[0], ACTIVITY_COLUMNS, activity_rows, ActivityRecord, problems, cell_refs)
    refrigerants = _parse(sheets[1], REFRIGERANT_COLUMNS, refrigerant_rows, RefrigerantRecord, problems, cell_refs)
    metrics = _parse_metrics(sheets[2], metric_rows, problems, cell_refs)
    _check_sources(config, activities, problems, cell_refs)
    if not activities and not problems:
        problems.append("There are no activity rows yet.")
    if problems:
        return None, problems
    return ActivityData(activities=activities, refrigerants=refrigerants, metrics=metrics), []


def read_input(path: str | Path, config: Config) -> ActivityData:
    path = Path(path)
    if path.is_dir():
        def rows(name):
            file = path / name
            if not file.exists():
                return []
            with file.open(newline="", encoding="utf-8-sig") as fh:
                reader = csv.reader(fh)
                next(reader, None)
                return [(i, row) for i, row in enumerate(reader, start=2)]
        activity_rows, refrigerant_rows, metric_rows = rows("activity.csv"), rows("refrigerants.csv"), rows("metrics.csv")
        sheets = ("activity.csv", "refrigerants.csv", "metrics.csv")
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        missing = [name for name in ("Activity", "Refrigerants", "Metrics") if name not in wb.sheetnames]
        if missing:
            raise InputError([f"{path.name}: missing sheet(s) {', '.join(missing)}; regenerate with make-template"])

        def rows(name):
            return [(i, row) for i, row in enumerate(wb[name].iter_rows(min_row=2, values_only=True), start=2)]
        activity_rows, refrigerant_rows, metric_rows = rows("Activity"), rows("Refrigerants"), rows("Metrics")
        sheets = ("Activity", "Refrigerants", "Metrics")

    data, problems = parse_tables(config, activity_rows, refrigerant_rows, metric_rows, sheets)
    if problems:
        raise InputError([p.replace("There are no activity rows yet.", f"{path.name}: no activity rows found")
                          for p in problems])
    return data

"""File outputs: line CSV, auditor workbook, Power BI tidy tables, run manifest."""

import csv
import hashlib
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ghg_inventory import __version__
from ghg_inventory.gwp import GWP_SOURCES, gwp
from ghg_inventory.inventory import InventoryResult, framework_mapping
from ghg_inventory.models import CATEGORY_LABELS, DATA_QUALITY

_GASES = ("CO2", "CH4", "N2O", "HFC-32", "HFC-125", "HFC-134a", "HFC-143a")

LINE_COLUMNS = (
    "line_id", "row_ref", "source_id", "source_name", "scope", "category", "method", "month",
    "quantity", "unit", "quantity_in_factor_unit", "factor_unit", "factor_id", "factor_description",
    *[f"kg_{g}" for g in _GASES], "kg_co2e", "t_co2e", "biogenic_co2_kg",
    "data_quality", "evidence_ref", "gwp_set", "data_status",
)


def _line_rows(result: InventoryResult) -> list[list]:
    rows = []
    for l in result.lines:
        rows.append([
            l.line_id, l.row_ref, l.source_id, l.source_name, l.scope, l.category, l.method, l.month,
            l.quantity, l.unit, round(l.quantity_in_factor_unit, 6), l.factor_unit, l.factor_id, l.factor_description,
            *[round(l.gas_kg.get(g, 0.0), 6) for g in _GASES], round(l.co2e_kg, 4), round(l.co2e_t, 6),
            round(l.biogenic_co2_kg, 4), l.data_quality, l.evidence_ref,
            result.config.reporting.gwp_set, result.config.reporting.data_status,
        ])
    return rows


def write_lines_csv(result: InventoryResult, path: Path) -> Path:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(LINE_COLUMNS)
        writer.writerows(_line_rows(result))
    return path


def write_tidy_tables(result: InventoryResult, folder: Path) -> list[Path]:
    """Star schema for Power BI / Tableau: one fact table, three dimensions, joined on ids."""
    folder.mkdir(parents=True, exist_ok=True)
    gwp_set = result.config.reporting.gwp_set
    tables = {
        "fact_emissions.csv": (
            ("line_id", "row_ref", "source_id", "factor_id", "scope", "method", "month", "gas",
             "kg_gas", "gwp", "kg_co2e", "gwp_set", "reporting_year"),
            [[l.line_id, l.row_ref, l.source_id, l.factor_id, l.scope, l.method, l.month, gas,
              kg, gwp(gas, gwp_set), kg * gwp(gas, gwp_set), gwp_set, result.config.reporting.year]
             for l in result.lines for gas, kg in l.gas_kg.items()],
        ),
        "dim_source.csv": (
            ("source_id", "source_name", "scope", "category", "category_label"),
            sorted({(l.source_id, l.source_name, l.scope, l.category, CATEGORY_LABELS[l.category])
                    for l in result.lines}),
        ),
        "dim_factor.csv": (
            ("factor_id", "library_id", "description", "unit", "co2e_by_gas", "gwp_basis",
             "uncertainty_pct", "publisher", "title", "year", "reference", "url"),
            [[f.factor_id, f.library.library_id, f.entry.description, f.unit, json.dumps(f.entry.co2e_by_gas),
              f.library.gwp_basis, f.entry.uncertainty_pct, f.library.publisher, f.library.title,
              f.library.year, f.entry.reference, f.library.url] for f in result.used_factors()],
        ),
        "dim_activity.csv": (
            ("row_ref", "source_id", "month", "quantity", "unit", "data_quality", "data_quality_label",
             "evidence_ref", "notes"),
            [[a.row_ref, a.source_id, a.month, a.quantity, a.unit, a.data_quality,
              DATA_QUALITY[a.data_quality][0], a.evidence_ref, a.notes] for a in result.data.activities]
            + [[r.row_ref, f"refrigerant_{r.equipment.lower().replace(' ', '_')}", str(result.config.reporting.year),
                r.released_kg, "kg", r.data_quality, DATA_QUALITY[r.data_quality][0], r.evidence_ref,
                f"{r.refrigerant}, {r.method}"] for r in result.data.refrigerants],
        ),
    }
    paths = []
    for name, (header, rows) in tables.items():
        path = folder / name
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow(header)
            writer.writerows(rows)
        paths.append(path)
    return paths


# ---------------------------------------------------------------- workbook

_HEADER_FILL = PatternFill("solid", fgColor="1F3A4D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _table(wb: Workbook, title: str, header: tuple | list, rows: list, widths: dict[int, int] | None = None):
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for idx in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = (widths or {}).get(idx, 16)
    ws.freeze_panes = "A2"
    return ws


def write_workbook(result: InventoryResult, path: Path) -> Path:
    cfg, t = result.config, result.totals
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    rows = [
        ("Organisation", cfg.organisation.name),
        ("Site", cfg.organisation.site),
        ("Reporting year", cfg.reporting.year),
        ("Base year", cfg.reporting.base_year),
        ("Consolidation approach", cfg.reporting.consolidation_approach),
        ("Data status", cfg.reporting.data_status),
        ("GWP set", f"{cfg.reporting.gwp_set}: {GWP_SOURCES[cfg.reporting.gwp_set]}"),
        ("", ""),
        ("Scope 1 (t CO2e)", round(t.scope1_kg / 1000, 3)),
        *[(f"  {CATEGORY_LABELS[c]}", round(kg / 1000, 3)) for c, kg in t.scope1_by_category_kg.items()],
        ("Scope 2, location-based (t CO2e)", round(t.scope2_location_kg / 1000, 3)),
        ("Scope 2, market-based (t CO2e)", round(t.scope2_market_kg / 1000, 3)),
        ("Total, location-based (t CO2e)", round(t.total_location_kg / 1000, 3)),
        ("Total, market-based (t CO2e)", round(t.total_market_kg / 1000, 3)),
        ("Biogenic CO2, outside scopes (t)", round(t.biogenic_co2_kg / 1000, 3)),
        ("Emissions-weighted data quality (1 best, 5 worst)", round(t.weighted_data_quality, 2)),
    ]
    for row in rows:
        summary.append(list(row))
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 60
    for cell in summary["A"]:
        cell.font = Font(bold=True)

    _table(wb, "Lines", LINE_COLUMNS, _line_rows(result), {1: 26, 4: 32, 14: 50})
    _table(wb, "Activity", ("row_ref", "source_id", "month", "quantity", "unit", "data_quality", "evidence_ref", "notes"),
           [[a.row_ref, a.source_id, a.month, a.quantity, a.unit, a.data_quality, a.evidence_ref, a.notes]
            for a in result.data.activities], {7: 26, 8: 36})
    _table(wb, "Refrigerants", ("row_ref", "equipment", "refrigerant", "charge_kg", "method", "annual_leak_rate",
                                "recharged_kg", "released_kg", "gwp", "t_co2e", "data_quality", "evidence_ref"),
           [[r.row_ref, r.equipment, r.refrigerant, r.charge_kg, r.method, r.annual_leak_rate, r.recharged_kg,
             r.released_kg, gwp(r.refrigerant, cfg.reporting.gwp_set),
             r.released_kg * gwp(r.refrigerant, cfg.reporting.gwp_set) / 1000, r.data_quality, r.evidence_ref]
            for r in result.data.refrigerants], {2: 32, 12: 40})
    _table(wb, "Factors", ("factor_id", "description", "unit", "co2e_by_gas (library GWP basis)", "gwp_basis",
                           "uncertainty_pct", "source", "reference", "url", "notes"),
           [[f.factor_id, f.entry.description, f.unit, json.dumps(f.entry.co2e_by_gas), f.library.gwp_basis,
             f.entry.uncertainty_pct, f"{f.library.publisher}, {f.library.title}", f.entry.reference,
             f.library.url, f.entry.notes] for f in result.used_factors()],
           {1: 32, 2: 40, 4: 40, 7: 50, 8: 60, 9: 50, 10: 60})
    _table(wb, "Energy", ("source_id", "source", "MWh", "renewable MWh", "renewable share"),
           [[e.source_id, e.source_name, e.kwh / 1000, e.renewable_kwh / 1000,
             e.renewable_kwh / e.kwh if e.kwh else 0] for e in result.energy], {2: 36})
    _table(wb, "Uncertainty", ("total", "central t", "mean t", "2.5th pct t", "97.5th pct t", "± % (95%)"),
           [[k, u.central_kg / 1000, u.mean_kg / 1000, u.low_kg / 1000, u.high_kg / 1000, u.half_width_pct]
            for k, u in result.uncertainty.items()], {1: 20})
    _table(wb, "Intensity", ("metric", "denominator", "location-based", "market-based"),
           [[i.label, i.denominator, i.location_based, i.market_based] for i in result.intensities], {1: 44})
    _table(wb, "Abatement", ("option", "type", "t saved/yr (LB)", "t saved/yr (MB)", "capex EUR",
                             "energy saving EUR/yr", "net annual cost EUR", "EUR per t (MB)", "payback years", "note"),
           [[a.name, a.type, a.saving_location_t, a.saving_market_t, a.capex_eur, a.annual_energy_saving_eur,
             a.annual_cost_eur, a.eur_per_t, a.payback_years, a.note] for a in result.abatement],
           {1: 44, 2: 24, 10: 50})
    scope3_rows = [[s.name, s.quantity, s.unit, s.factor_id, s.estimate_t, s.multiple_of_scope12_market, s.citation]
                   for s in result.scope3]
    _table(wb, "Scope 3 screen", ("item (indicative, outside boundary)", "quantity", "unit", "factor_id",
                                  "t CO2e", "x Scope 1+2 (MB)", "source"), scope3_rows, {1: 40, 7: 80})
    _table(wb, "Framework mapping", ("framework", "disclosure", "value"), framework_mapping(result), {1: 34, 2: 60, 3: 70})
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.000"
    wb.save(path)
    return path


# ---------------------------------------------------------------- manifest


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _git_commit(folder: Path) -> str | None:
    try:
        out = subprocess.run(["git", "rev-parse", "HEAD"], cwd=folder, capture_output=True, text=True, timeout=5)
    except (OSError, subprocess.SubprocessError):
        return None
    return out.stdout.strip() if out.returncode == 0 else None


def write_manifest(result: InventoryResult, config_path: Path, input_path: Path, base_dir: Path,
                   outputs: list[Path], path: Path) -> Path:
    cfg = result.config
    manifest = {
        "tool": "ghg-inventory",
        "tool_version": __version__,
        "run_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "git_commit": _git_commit(base_dir),
        "organisation": cfg.organisation.name,
        "reporting_year": cfg.reporting.year,
        "data_status": cfg.reporting.data_status,
        "gwp_set": cfg.reporting.gwp_set,
        "inputs": {
            "config": {"path": str(config_path), "sha256": _sha256(config_path)},
            "activity_data": {"path": str(input_path),
                              "sha256": _sha256(input_path) if input_path.is_file() else None},
            "factor_libraries": [
                {"path": p, "sha256": _sha256(base_dir / p)} for p in cfg.factor_libraries
            ],
        },
        "totals_t_co2e": {
            "scope1": round(result.totals.scope1_kg / 1000, 3),
            "scope2_location": round(result.totals.scope2_location_kg / 1000, 3),
            "scope2_market": round(result.totals.scope2_market_kg / 1000, 3),
            "total_location": round(result.totals.total_location_kg / 1000, 3),
            "total_market": round(result.totals.total_market_kg / 1000, 3),
        },
        "outputs": [str(p) for p in outputs],
    }
    path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    return path

import csv
import json

import pytest
from openpyxl import load_workbook

from ghg_inventory.cli import main
from ghg_inventory.excel_io import read_input
from ghg_inventory.inventory import framework_mapping, run_inventory
from ghg_inventory.report import build_report_context


@pytest.fixture(scope="module")
def result(loaded):
    config, factors, base_dir = loaded
    return run_inventory(config, factors, read_input(base_dir / config.input, config))


def test_representative_inventory_shape(result):
    t = result.totals
    assert t.scope2_location_kg == pytest.approx(200000 * 0.344)
    assert t.scope2_location_kg > t.scope1_kg            # electricity dominates, as expected for a cutter
    assert t.scope2_market_kg > t.scope2_location_kg     # residual mix is dirtier than the grid average


def test_report_context(result):
    c = build_report_context(result)
    assert c["representative"] is True
    assert c["largest_name"] == "Purchased grid electricity"
    assert sum(g["location_kg"] for g in c["line_items"]) == pytest.approx(c["total_location_kg"])
    assert sum(g["market_kg"] for g in c["line_items"]) == pytest.approx(c["total_market_kg"])
    assert any(a.type == "heat_pump" for a in c["backfires"])
    assert c["scope3"][0].multiple_of_scope12_market > 1


def test_framework_mapping_covers_vsme_esrs_iso(result):
    frameworks = {row[0] for row in framework_mapping(result)}
    assert {"VSME (EFRAG) B3", "ESRS E1-6", "ISO 14064-1:2018"} <= frameworks


def test_cli_writes_everything(tmp_path, loaded):
    _, _, base_dir = loaded
    code = main(["run", "--config", str(base_dir / "config.yaml"), "--output-dir", str(tmp_path)])
    assert code == 0
    for name in ("inventory_lines.csv", "ghg_inventory.xlsx", "ghg_inventory_report.pdf", "run_manifest.json",
                 "tables/fact_emissions.csv", "tables/dim_source.csv", "tables/dim_factor.csv",
                 "tables/dim_activity.csv"):
        assert (tmp_path / name).stat().st_size > 0, name
    assert (tmp_path / "ghg_inventory_report.pdf").read_bytes()[:4] == b"%PDF"

    with (tmp_path / "inventory_lines.csv").open(encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    location = sum(float(r["kg_co2e"]) for r in rows if r["method"] in ("direct", "location"))
    with (tmp_path / "tables" / "fact_emissions.csv").open(encoding="utf-8") as fh:
        fact = list(csv.DictReader(fh))
    fact_location = sum(float(r["kg_co2e"]) for r in fact if r["method"] in ("direct", "location"))
    assert fact_location == pytest.approx(location, rel=1e-6)    # per-gas fact table reconciles to lines

    manifest = json.loads((tmp_path / "run_manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["inputs"]["config"]["sha256"]) == 64
    assert manifest["totals_t_co2e"]["scope2_location"] == pytest.approx(68.8)

    wb = load_workbook(tmp_path / "ghg_inventory.xlsx", read_only=True)
    assert {"Summary", "Lines", "Factors", "Uncertainty", "Abatement", "Framework mapping"} <= set(wb.sheetnames)


def test_cli_bad_config_returns_2(tmp_path, capsys):
    bad = tmp_path / "config.yaml"
    bad.write_text("organisation: [", encoding="utf-8")
    assert main(["run", "--config", str(bad)]) == 2
    assert "not valid YAML" in capsys.readouterr().err


def test_cli_make_template(tmp_path, loaded):
    _, _, base_dir = loaded
    out = tmp_path / "t.xlsx"
    assert main(["make-template", "--config", str(base_dir / "config.yaml"), "--out", str(out)]) == 0
    assert "Activity" in load_workbook(out).sheetnames

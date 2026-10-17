import pytest
from openpyxl import load_workbook

from conftest import activity, data, refrigerant
from ghg_inventory.excel_io import InputError, make_template, read_input


def test_blank_template_has_sheets_and_dropdowns(loaded, tmp_path):
    config, factors, _ = loaded
    path = make_template(config, factors, tmp_path / "t.xlsx")
    wb = load_workbook(path)
    assert {"Instructions", "Activity", "Refrigerants", "Metrics", "Sources", "Lists"} <= set(wb.sheetnames)
    assert wb["Lists"].sheet_state == "hidden"
    assert wb["Activity"].protection.sheet is True
    assert wb["Activity"]["A2"].protection.locked is False          # data cells editable
    ranges = [str(dv.sqref) for dv in wb["Activity"].data_validations.dataValidation]
    assert any(r.startswith("A2") for r in ranges) and any(r.startswith("E2") for r in ranges)
    sources = {row[0]: row[5] for row in wb["Sources"].iter_rows(min_row=2, values_only=True)}
    assert sources["vehicle_diesel"] == "litre" and sources["grid_electricity"] == "kWh"


def test_round_trip(loaded, tmp_path):
    config, factors, _ = loaded
    original = data(
        activity("grid_electricity", 1234.5, "kWh", row="x"),
        activity("vehicle_diesel", 50, "litre", month="2025-03", dq=2, row="y"),
        refrigerants=[refrigerant("R-410A", charge=2, leak=0.05, row="z")],
        metrics={"fte": 12},
    )
    path = make_template(config, factors, tmp_path / "filled.xlsx", data=original)
    back = read_input(path, config)
    assert [(a.source_id, a.month, a.quantity, a.data_quality) for a in back.activities] == [
        ("grid_electricity", "2025-01", 1234.5, 1), ("vehicle_diesel", "2025-03", 50, 2)]
    assert back.activities[0].row_ref == "Activity!2"
    assert back.refrigerants[0].released_kg == pytest.approx(0.1)
    assert back.metrics == {"fte": 12}


def _put(ws, rows):
    """Write rows from row 2 (the template's formatted cells make ws.append start at row 601)."""
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)


def _filled(loaded, tmp_path, rows):
    config, factors, _ = loaded
    path = make_template(config, factors, tmp_path / "bad.xlsx")
    wb = load_workbook(path)
    _put(wb["Activity"], rows)
    wb.save(path)
    return path, config


def test_errors_point_at_cells(loaded, tmp_path):
    path, config = _filled(loaded, tmp_path, [
        ["grid_electricity", "2025-01", -5, "kWh", 1],
        ["heating_gas", "2025-13", 100, "kWh", 1],
        ["heating_gas", "2025-02", 100, "kWh", 9],
        ["steam", "2025-02", 100, "kWh", 1],
        ["heating_gas", "2024-12", 100, "kWh", 1],
    ])
    with pytest.raises(InputError) as exc:
        read_input(path, config)
    text = "\n".join(exc.value.problems)
    assert "Activity!C2 (quantity)" in text
    assert "Activity!B3 (month)" in text
    assert "Activity!E4 (data_quality)" in text
    assert "Activity!5 (source_id)" in text and "steam" in text
    assert "outside the reporting year" in text


def test_unknown_refrigerant_and_missing_leak_rate(loaded, tmp_path):
    config, factors, _ = loaded
    path = make_template(config, factors, tmp_path / "r.xlsx")
    wb = load_workbook(path)
    _put(wb["Activity"], [["grid_electricity", "2025-01", 1, "kWh", 1]])
    _put(wb["Refrigerants"], [["AC", "R-999", 2, "screening", 0.05, None, 4],
                              ["AC2", "R-32", 2, "screening", None, None, 4]])
    wb.save(path)
    with pytest.raises(InputError) as exc:
        read_input(path, config)
    text = "\n".join(exc.value.problems)
    assert "Refrigerants!B2 (refrigerant)" in text
    assert "Refrigerants!3" in text and "annual_leak_rate" in text


def test_csv_folder_input(loaded, tmp_path):
    config, _, _ = loaded
    (tmp_path / "activity.csv").write_text(
        "source_id,month,quantity,unit,data_quality,evidence_ref,notes\n"
        "heating_gas,2025-01,1000,m3,1,INV-1,\n", encoding="utf-8")
    (tmp_path / "metrics.csv").write_text("metric,value\nfte,20\n", encoding="utf-8")
    d = read_input(tmp_path, config)
    assert d.activities[0].row_ref == "activity.csv!2"
    assert d.activities[0].unit == "m3"
    assert d.metrics == {"fte": 20}
    assert d.refrigerants == []


def test_project_representative_workbook_reads(loaded):
    config, _, base_dir = loaded
    d = read_input(base_dir / config.input, config)
    assert len(d.activities) == 40
    totals = {}
    for a in d.activities:
        totals[a.source_id] = totals.get(a.source_id, 0) + a.quantity
    assert totals == {"grid_electricity": 200000, "heating_gas": 180000, "vehicle_diesel": 2000, "forklift_lpg": 1000}
